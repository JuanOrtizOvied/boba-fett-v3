"""Client-template portfolio export.

Fills the "ACTUAL (* NO TOCAR *)" column (D) of the attached client template
(`templates/portafolio_final_template.xlsx`) from the user's portfolio, by
taxonomy leaf, as a fraction of the total portfolio.

Asset-class level data is reliable; `underlying` is free text. So only
amounts that match a leaf with confidence are written. Everything else is
reported on a second "Sin clasificar (SABBI)" sheet and never guessed into a
leaf.
"""

from __future__ import annotations

import io
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from db.excel import ASSET_CLASS_LABELS, ASSET_CLASS_ORDER, HEADER_FILL, HEADER_FONT
from db.models import Product

TEMPLATE_PATH = Path(__file__).parent / "templates" / "portafolio_final_template.xlsx"
NOTE_SHEET_TITLE = "Sin clasificar (SABBI)"
PERCENT_FORMAT = "0.0%"
CASH_CELL = "D89"

# (cell, asset_class_key, group, leaf) — kept in sync with `ASSET_CLASSES` in
# agent/state.py. Cash is a single cell fed by the whole class (see below).
LEAF_CELLS: list[tuple[str, str, str, str]] = [
    ("D17", "inversiones_directas", "RE Perú", "Residencial"),
    ("D18", "inversiones_directas", "RE Perú", "Oficinas"),
    ("D19", "inversiones_directas", "RE Perú", "Comercial/Industrial"),
    ("D21", "inversiones_directas", "RE Extranjero", "RE Extranjero"),
    ("D28", "mercados_privados", "Deuda Privada", "Deuda Privada"),
    ("D30", "mercados_privados", "Private Equity", "Private Equity"),
    ("D32", "mercados_privados", "Venture Capital", "Venture Capital"),
    ("D34", "mercados_privados", "Real Estate", "Real Estate"),
    ("D36", "mercados_privados", "Hedge Funds", "Hedge Funds"),
    ("D38", "mercados_privados", "Infraestructura", "Infraestructura"),
    ("D45", "club_deals", "Real Estate", "Perú"),
    ("D46", "club_deals", "Real Estate", "Extranjero"),
    ("D48", "club_deals", "Deuda Privada", "Perú"),
    ("D49", "club_deals", "Deuda Privada", "Extranjero"),
    ("D51", "club_deals", "Otros", "Perú"),
    ("D52", "club_deals", "Otros", "Extranjero"),
    ("D59", "mercados_publicos", "Renta Variable", "US Large Cap"),
    ("D60", "mercados_publicos", "Renta Variable", "US Mid & Small Cap"),
    ("D61", "mercados_publicos", "Renta Variable", "Developed ex-US"),
    ("D62", "mercados_publicos", "Renta Variable", "EM ex-Perú"),
    ("D63", "mercados_publicos", "Renta Variable", "Perú"),
    ("D65", "mercados_publicos", "Renta Fija", "US Treasuries"),
    ("D66", "mercados_publicos", "Renta Fija", "IG Corporates AAA-BBB"),
    ("D67", "mercados_publicos", "Renta Fija", "High Yield BB-"),
    ("D68", "mercados_publicos", "Renta Fija", "EM Bonds"),
    ("D69", "mercados_publicos", "Renta Fija", "LatAm Bonds"),
    ("D70", "mercados_publicos", "Renta Fija", "Perú Bonds"),
    ("D77", "otros", "Cripto", "Bitcoin"),
    ("D78", "otros", "Cripto", "Ethereum"),
    ("D79", "otros", "Cripto", "Otras"),
    ("D81", "otros", "Commodities", "Oro"),
    ("D82", "otros", "Commodities", "Otros"),
]

# Same resolution as `resolveAssetClassKey` in apps/web/lib/categories.ts.
_CLASS_ALIASES: dict[str, str] = {
    **{label.lower(): key for key, label in ASSET_CLASS_LABELS.items()},
    "real estate directo": "inversiones_directas",
    "mercados privado": "mercados_privados",
    "directas": "inversiones_directas",
    "privados": "mercados_privados",
    "club": "club_deals",
    "publicos": "mercados_publicos",
    "cash": "cash_y_equivalentes",
}

# Conservative abbreviation prefixes seen in agent output.
_PREFIX_ALIASES: dict[str, str] = {"rv ": "renta variable ", "rf ": "renta fija "}


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.strip().casefold())
    return " ".join("".join(c for c in decomposed if not unicodedata.combining(c)).split())


def _resolve_class(name: str) -> str:
    key = name.strip()
    if key in ASSET_CLASS_ORDER:
        return key
    return _CLASS_ALIASES.get(key.lower(), "otros")


def _match_leaf(class_key: str, underlying: str) -> str | None:
    """Return the cell for `underlying` within the class, or None when there is
    no unique confident match. Ambiguity at one step falls through to the next."""
    candidates = [c for c in LEAF_CELLS if c[1] == class_key]
    norm = _norm(underlying)
    variants = [norm]
    for prefix, full in _PREFIX_ALIASES.items():
        if norm.startswith(prefix):
            variants.append(full + norm[len(prefix) :])
    for text in variants:
        for forms in (
            lambda c: _norm(c[3]),
            lambda c: _norm(f"{c[2]} {c[3]}"),
            lambda c: _norm(c[2]),
        ):
            hits = [c[0] for c in candidates if forms(c) == text]
            if len(hits) == 1:
                return hits[0]
    return None


def _write_note_sheet(wb, rows: list[list]) -> None:
    ws = wb.create_sheet(title=NOTE_SHEET_TITLE)
    ws["A1"] = (
        "Estos montos no pudieron clasificarse en una subcategoría con certeza y "
        "quedaron FUERA de la columna ACTUAL. El % total por clase es confiable."
    )
    ws["A1"].font = Font(italic=True)
    headers = ["Clase de activo", "% mapeado", "% sin clasificar", "% total clase", "Detalle"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=4):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if 2 <= col <= 4:
                cell.number_format = PERCENT_FORMAT
    for letter, width in zip("ABCDE", (24, 14, 18, 16, 80)):
        ws.column_dimensions[letter].width = width


def build_portfolio_template_workbook(products: list[Product]) -> io.BytesIO:
    """Load the client template fresh, fill ACTUAL (column D) and append the
    "Sin clasificar (SABBI)" sheet. The on-disk asset is never modified."""
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.worksheets[0]

    total = sum(p.amount for p in products)
    cell_amounts: dict[str, float] = {}
    mapped: dict[str, float] = {}
    unmapped: dict[str, float] = {}
    details: dict[str, list[str]] = {}

    def note(class_key: str, text: str) -> None:
        if text not in details.setdefault(class_key, []):
            details[class_key].append(text)

    if total > 0:
        for p in products:
            for ac in p.asset_class:
                class_key = _resolve_class(ac.name)
                class_amount = p.amount * ac.percentage / 100
                if class_key == "cash_y_equivalentes":
                    cell_amounts[CASH_CELL] = cell_amounts.get(CASH_CELL, 0.0) + class_amount
                    mapped[class_key] = mapped.get(class_key, 0.0) + class_amount
                    continue
                mapped_here = 0.0
                for u in p.underlying:
                    amount = class_amount * u.percentage / 100
                    cell = _match_leaf(class_key, u.name)
                    if cell is None:
                        note(class_key, f"{p.name} — {u.name}")
                        continue
                    cell_amounts[cell] = cell_amounts.get(cell, 0.0) + amount
                    mapped_here += amount
                rest = max(class_amount - mapped_here, 0.0)
                if not p.underlying:
                    note(class_key, f"{p.name} — (sin subyacente)")
                mapped[class_key] = mapped.get(class_key, 0.0) + mapped_here
                unmapped[class_key] = unmapped.get(class_key, 0.0) + rest

        for cell, amount in cell_amounts.items():
            ws[cell] = amount / total

    rows: list[list] = []
    if total > 0:
        for class_key in ASSET_CLASS_ORDER:
            m, u = mapped.get(class_key, 0.0), unmapped.get(class_key, 0.0)
            if m + u <= 0:
                continue
            rows.append(
                [
                    ASSET_CLASS_LABELS[class_key],
                    m / total,
                    u / total,
                    (m + u) / total,
                    "; ".join(details.get(class_key, [])),
                ]
            )
    _write_note_sheet(wb, rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def template_export_filename() -> str:
    return f"portafolio-sabbi-{date.today().isoformat()}.xlsx"
