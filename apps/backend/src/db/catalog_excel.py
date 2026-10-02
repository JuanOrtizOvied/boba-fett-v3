"""Server-side .xlsx export for selected `product_catalog` entries.

Builds a flat, single-sheet listing workbook directly from Postgres data —
no client-side spreadsheet library, no file ever written to disk. Kept
separate from `db/excel.py` (the portfolio export) because that module is
shaped around one user's holdings, proportionally split across asset
classes, while a catalog export is a flat listing of independent entries
with no per-user context (see
`openspec/changes/catalog-export-and-filters/design.md` ADR-2).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from catalog_sync.config import CONTROL_FIELD_TYPE, FIELD_MAPPING, SHEET_NAME
from db.models import AssetAllocation, CatalogProduct

HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

# Mirrors `CATALOG_COLUMNS` in `apps/web/app/admin/catalog/page.tsx` — keep
# both lists in sync so the exported file matches what admins see in the
# table. `id` is exported first even though it isn't a visible table
# column, so a future re-import can target existing rows for update
# (design.md ADR-3).
_COLUMNS: list[tuple[str, str]] = [
    ("id", "ID"),
    ("codigo", "Codigo"),
    ("name", "Nombre"),
    ("alternative_names", "Nombres alternativos"),
    ("asset_class", "Clase de activo"),
    ("geographic_focus", "Foco geográfico"),
    ("underlying", "Subyacente"),
    ("commission", "Comisión"),
    ("currency", "Moneda"),
    ("administrator", "Administrador"),
    ("manager", "Gestor"),
    ("liquidity", "Liquidez"),
    ("return_rate", "Rentabilidad"),
    ("cash_flows", "Flujos"),
    ("isin", "ISIN"),
    ("distribution", "Distribución"),
]


def _style_header_row(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize_columns(ws, count: int, width: int = 26) -> None:
    for col_idx in range(1, count + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _format_allocations(allocations: list[AssetAllocation]) -> str:
    """Renders `[{name, percentage}]` the same way the admin catalog table
    already does — e.g. "Renta Fija 60%, Renta Variable 40%"."""
    return ", ".join(f"{a.name} {a.percentage:.0f}%" for a in allocations)


def _format_cell(entry: CatalogProduct, field: str) -> object:
    value = getattr(entry, field)
    if field in ("asset_class", "geographic_focus", "underlying"):
        return _format_allocations(value)
    if field == "alternative_names":
        return ", ".join(value)
    return value


def build_catalog_workbook(entries: list[CatalogProduct]) -> io.BytesIO:
    """Build a single-sheet workbook listing the given catalog entries, one
    row per entry, and return it as an in-memory buffer ready for
    streaming. No intermediate file is written to disk."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Catálogo"

    headers = [label for _, label in _COLUMNS]
    _style_header_row(ws, headers)

    for row_idx, entry in enumerate(entries, start=2):
        for col_idx, (field, _label) in enumerate(_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_format_cell(entry, field))

    _autosize_columns(ws, len(headers))
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_filename() -> str:
    return f"catalogo-sabbi-{date.today().isoformat()}.xlsx"


def export_filename_for_sync() -> str:
    return f"catalogo-sabbi-sharepoint-{date.today().isoformat()}.xlsx"


@dataclass(frozen=True)
class SyncWorkbook:
    """A workbook in the layout the SharePoint sync reads, plus what was left
    out of it: the sync matches rows by `codigo`, so an entry without one has
    no place in it (`skipped_without_codigo` holds their names)."""

    buffer: io.BytesIO
    included: int
    skipped_without_codigo: tuple[str, ...]


def _format_percentage(value: float) -> str:
    """The exact value with a dot decimal and no trailing zeros ("85", "33.3"),
    so reading it back gives the same number. No exponent: the composite
    parser only accepts digits and a dot."""
    value = float(value)
    if value.is_integer():
        return str(int(value))
    text = repr(value)
    return text if "e" not in text.lower() else format(value, "f").rstrip("0").rstrip(".")


def _format_composite(allocations: list[AssetAllocation]) -> str:
    return ", ".join(f"{a.name} {_format_percentage(a.percentage)}%" for a in allocations)


def _sync_cell(entry: CatalogProduct, spec: dict) -> object:
    if spec["type"] == CONTROL_FIELD_TYPE:
        return None  # `Eliminar`: a product is never deleted by being in the file
    value = getattr(entry, spec["db_field"])
    if spec["type"] == "composite":
        return _format_composite(value) or None
    return value or None  # an empty value is an empty cell, not ""


def build_sync_workbook(entries: list[CatalogProduct]) -> SyncWorkbook:
    """Build the workbook the SharePoint sync reads: sheet `SHEET_NAME`, one
    header per `FIELD_MAPPING` key in its order, composites as
    "Name 85%, Name 15%" with exact decimals, and `Eliminar` empty. Because
    the headers come from the mapping, a column added there appears here with
    no further change. Soft-deleted entries are left out silently; entries
    without a `codigo` are left out and reported."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    headers = list(FIELD_MAPPING)
    _style_header_row(ws, headers)

    skipped: list[str] = []
    row_idx = 2
    for entry in entries:
        if entry.is_deleted:
            continue
        if not (entry.codigo or "").strip():
            skipped.append(entry.name)
            continue
        for col_idx, spec in enumerate(FIELD_MAPPING.values(), start=1):
            ws.cell(row=row_idx, column=col_idx, value=_sync_cell(entry, spec))
        row_idx += 1

    _autosize_columns(ws, len(headers))
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return SyncWorkbook(
        buffer=buffer, included=row_idx - 2, skipped_without_codigo=tuple(skipped)
    )
