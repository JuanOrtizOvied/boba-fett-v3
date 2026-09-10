"""Ficha Patrimonial Excel parser and Pydantic contracts.

Parses the admin-uploaded "Ficha Patrimonial" `.xlsx` template (sheet
"Sabbi") into structured rows the admin ficha-import flow can enrich via
`cascade_search()` and review before bulk-creating products
(`sdd/admin-ficha-patrimonial/spec` — "Excel Parsing — Sabbi Sheet",
"Return Rate Normalization", "Asset Class Mapping").

Column/row layout is centralized here as named constants — never hardcode
these literals elsewhere (spec: "Column mapping is centralized").
"""

from __future__ import annotations

import io
import unicodedata
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from db.models import AssetAllocation, SearchResult

# --- Sheet layout constants ---------------------------------------------

SHEET_NAME = "Sabbi"

CLIENT_INFO_ROW = 5
COL_CLIENT_NAME = "D"
COL_CLIENT_EMAIL = "M"
COL_CLIENT_PHONE = "O"

HEADER_ROW = 7
PRODUCTS_START_ROW = 8
# Real-estate rows start here and are always excluded from parsed products
# (spec: "Real estate rows excluded").
REAL_ESTATE_START_ROW = 27

COL_PRODUCT_NAME = "C"
COL_TIPO_ACTIVO = "J"
COL_CURRENCY = "K"
COL_PERTENENCIA = "L"
COL_AMOUNT = "M"
COL_RETURN_RATE = "N"


class FichaParseError(ValueError):
    """Raised when an uploaded workbook cannot be parsed as a valid ficha
    patrimonial: corrupt/non-`.xlsx` file or missing the "Sabbi" sheet
    (spec: "Invalid or corrupt file rejected", "Missing Sabbi sheet")."""


# --- Tipo de activo -> asset_class mapping ------------------------------

# Hardcoded, finite lookup (spec: "Asset Class Mapping" — a fixed table, no
# guessing). Keys are normalized (lowercased, unaccented, whitespace
# collapsed) via `_normalize_lookup_key` before matching, so the source
# Excel free-text casing/accents don't matter.
TIPO_ACTIVO_MAP: dict[str, str] = {
    "acciones en bolsa": "mercados_publicos",
    "fondos mutuos": "mercados_publicos",
    "cuenta ahorros": "cash_y_equivalentes",
    "cuenta corriente": "cash_y_equivalentes",
    "deposito plazo fijo": "cash_y_equivalentes",
    "inversiones alternativas (private equity/venture capital/etc)": "mercados_privados",
}


def _normalize_lookup_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.strip().lower())
    without_accents = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return " ".join(without_accents.split())


def map_asset_class(tipo_activo: str | None) -> str | None:
    """Map a free-text "Tipo de activo" value to a SABBI `ASSET_CLASSES`
    top-level key, or `None` when unrecognized (spec: "Unknown value
    flagged, not guessed" — including the literal "Otro", which the
    template uses as a catch-all and is intentionally NOT mapped so the
    admin must classify it manually)."""
    if not tipo_activo:
        return None
    return TIPO_ACTIVO_MAP.get(_normalize_lookup_key(tipo_activo))


# --- Return rate normalization -------------------------------------------


def normalize_return_rate(value: Any) -> str:
    """Normalize a return-rate cell into one decimal-string format
    (spec: "Return Rate Normalization"):

    - Percentage string ("7.5%") -> strip "%", format to 1 decimal ("7.5").
    - Bare numeric string ("8") -> format to 1 decimal ("8.0").
    - Fractional decimal (0.08, from a percent-formatted cell) -> multiply
      by 100, format to 1 decimal ("8.0").
    - Missing/blank/non-numeric -> "" (never raises).
    """
    if value is None:
        return ""
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return ""
        if stripped.endswith("%"):
            stripped = stripped[:-1].strip()
        try:
            number = float(stripped)
        except ValueError:
            return ""
        return f"{number:.1f}"
    if isinstance(value, (int, float)):
        return f"{value * 100:.1f}"
    return ""


# --- Pydantic contracts ----------------------------------------------------


class FichaParseRequest(BaseModel):
    """Request body for `POST /admin/ficha-patrimonial/parse`."""

    file_data: str = Field(description="Base64-encoded .xlsx content")
    file_name: str = ""


class FichaClientInfo(BaseModel):
    """Client info extracted from row 5. `user_id` is empty until the
    parse endpoint resolves it via `UserRepository.get_by_email()` (PR 2)."""

    email: str = ""
    name: str = ""
    phone: str = ""
    user_id: str = ""


class FichaParsedRow(BaseModel):
    """One raw product row (rows 8+), before enrichment."""

    excel_row: int
    raw_name: str
    raw_tipo_activo: str = ""
    raw_amount: float = 0
    raw_currency: str = ""
    raw_return_rate: str = ""
    raw_pertenencia: str = ""
    mapped_asset_class: str | None = None
    normalized_return_rate: str = ""


class ParsedFicha(BaseModel):
    """Return type of `parse_ficha_excel()` — client info + raw product
    rows, before `cascade_search()` enrichment (added in PR 2)."""

    client: FichaClientInfo
    rows: list[FichaParsedRow]


class FichaEnrichedRow(FichaParsedRow):
    """A parsed row merged with its `cascade_search()` enrichment result."""

    enriched: SearchResult | None = None
    enrichment_failed: bool = False


class FichaParseResponse(BaseModel):
    """Response body for `POST /admin/ficha-patrimonial/parse`."""

    client: FichaClientInfo
    rows: list[FichaEnrichedRow]


class FichaConfirmRow(BaseModel):
    """One admin-reviewed row submitted to `POST
    /admin/ficha-patrimonial/confirm`. Mirrors `ProductCreate`."""

    name: str
    provider: str = ""
    amount: float = Field(gt=0)
    underlying: list[AssetAllocation] = Field(default_factory=list)
    asset_class: list[AssetAllocation]
    geographic_focus: list[AssetAllocation] = Field(default_factory=list)
    commission: str = ""
    currency: str = ""
    administrator: str = ""
    manager: str = ""
    liquidity: str = ""
    return_rate: str = ""
    catalog_product_id: int | None = None


class FichaConfirmRequest(BaseModel):
    """Request body for `POST /admin/ficha-patrimonial/confirm`."""

    user_id: str
    products: list[FichaConfirmRow]


# --- Excel parsing -----------------------------------------------------


def _cell_str(ws: Any, row: int, col: str) -> str:
    value = ws[f"{col}{row}"].value
    if value is None:
        return ""
    return str(value).strip()


def _cell_float(ws: Any, row: int, col: str) -> float:
    value = ws[f"{col}{row}"].value
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def _is_total_row(name: str) -> bool:
    return "total" in name.strip().lower()


def _extract_client_info(ws: Any) -> FichaClientInfo:
    return FichaClientInfo(
        name=_cell_str(ws, CLIENT_INFO_ROW, COL_CLIENT_NAME),
        email=_cell_str(ws, CLIENT_INFO_ROW, COL_CLIENT_EMAIL),
        phone=_cell_str(ws, CLIENT_INFO_ROW, COL_CLIENT_PHONE),
    )


def _extract_product_rows(ws: Any) -> list[FichaParsedRow]:
    rows: list[FichaParsedRow] = []
    row_idx = PRODUCTS_START_ROW

    while row_idx < REAL_ESTATE_START_ROW:
        name = _cell_str(ws, row_idx, COL_PRODUCT_NAME)
        if not name or _is_total_row(name):
            break

        amount = _cell_float(ws, row_idx, COL_AMOUNT)
        if not amount:
            row_idx += 1
            continue

        tipo_activo = _cell_str(ws, row_idx, COL_TIPO_ACTIVO)
        raw_return_rate_value = ws[f"{COL_RETURN_RATE}{row_idx}"].value

        rows.append(
            FichaParsedRow(
                excel_row=row_idx,
                raw_name=name,
                raw_tipo_activo=tipo_activo,
                raw_amount=amount,
                raw_currency=_cell_str(ws, row_idx, COL_CURRENCY),
                raw_return_rate=(
                    "" if raw_return_rate_value is None else str(raw_return_rate_value)
                ),
                raw_pertenencia=_cell_str(ws, row_idx, COL_PERTENENCIA),
                mapped_asset_class=map_asset_class(tipo_activo),
                normalized_return_rate=normalize_return_rate(raw_return_rate_value),
            )
        )
        row_idx += 1

    return rows


def parse_ficha_excel(file_bytes: bytes) -> ParsedFicha:
    """Parse a ficha patrimonial `.xlsx` (sheet "Sabbi") into client info
    plus product rows. Raises `FichaParseError` on a corrupt/non-`.xlsx`
    file or a missing "Sabbi" sheet."""
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise FichaParseError(
            "No se pudo leer el archivo: verifica que sea un .xlsx válido"
        ) from exc

    if SHEET_NAME not in workbook.sheetnames:
        raise FichaParseError(f"No se encontró la hoja '{SHEET_NAME}' en el archivo")

    worksheet = workbook[SHEET_NAME]
    return ParsedFicha(
        client=_extract_client_info(worksheet),
        rows=_extract_product_rows(worksheet),
    )
