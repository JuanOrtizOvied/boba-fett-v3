"""Excel workbook parser for the normalized catalog: bytes -> typed rows of
the three sheets, keyed by DB field name.

Pure module: no network, no database. A row that cannot be trusted (blank
key, blank required value, a number that is not a number, a repeated key) is
skipped and reported instead of raising, so one bad row never blocks the
rest. A missing sheet or a missing required column does raise: the file is
not the workbook this sync understands and nothing should be applied.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from catalog_normalized_sync.config import SHEETS, Column, SheetSpec
from catalog_sync.normalize import normalize_key
from db.seed_catalog import _clean, _parse_pct_list


class MissingSheetError(ValueError):
    """A sheet this sync needs is not in the workbook."""


class MissingColumnError(ValueError):
    """A required column is not in the header row of its sheet."""


class _InvalidRowError(ValueError):
    """The row is skipped; the message is the reason shown in the report."""


@dataclass(frozen=True)
class ParsedRow:
    """One data row. `key` is the tuple of the sheet's key fields and
    `values` holds every synced column present in the sheet (keys included):
    text as `str`, numbers as `float | None`, percentage lists as a tuple of
    `(name, percentage)`."""

    row_number: int
    key: tuple[str, ...]
    values: dict[str, Any]


@dataclass(frozen=True)
class SkippedRow:
    """`key` is the row's business key when it could be read. The diff counts
    such a row as still in the workbook, so a row that only has a bad value
    never gets its database row soft-deleted."""

    sheet: str
    row_number: int
    reason: str
    key: tuple[str, ...] | None = None


@dataclass(frozen=True)
class ParsedWorkbook:
    products: tuple[ParsedRow, ...]
    series: tuple[ParsedRow, ...]
    admins: tuple[ParsedRow, ...]
    skipped: tuple[SkippedRow, ...]


def _text(value: object) -> str:
    """Cell as trimmed text; a whole float drops the trailing `.0`."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean(value)


def _number(value: object) -> float | None:
    """Blank is `None`. Text that is not a plain number (for example `3.84%`)
    is rejected instead of guessed: the DB stores fractions."""
    if isinstance(value, bool):
        raise ValueError("boolean")
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean(value)
    if not text:
        return None
    return float(text.replace(",", "."))


def _allocations(value: object, *, fallback_100: bool) -> tuple[tuple[str, float], ...]:
    items = json.loads(_parse_pct_list(_text(value), fallback_100=fallback_100))
    return tuple((item["name"], float(item["percentage"])) for item in items)


def _convert(column: Column, raw: object) -> Any:
    if column.type == "number":
        try:
            return _number(raw)
        except ValueError:
            raise _InvalidRowError(f"non-numeric value in {column.headers[0]!r}") from None
    if column.type == "allocations":
        return _allocations(raw, fallback_100=column.fallback_100)
    return _text(raw)


def _header_index(header_row: tuple, spec: SheetSpec) -> dict[str, int]:
    """DB field -> 0-based column position, matching headers by their
    normalized text (case, accents and spacing do not matter)."""
    lookup: dict[str, int] = {}
    for position, cell in enumerate(header_row):
        if cell not in (None, ""):
            lookup.setdefault(normalize_key(str(cell)), position)

    index: dict[str, int] = {}
    missing: list[str] = []
    for column in spec.columns:
        position = next(
            (lookup[key] for key in map(normalize_key, column.headers) if key in lookup), None
        )
        if position is not None:
            index[column.field] = position
        elif column.required:
            missing.append(column.headers[0])
    if missing:
        raise MissingColumnError(f"Sheet {spec.sheet_name!r} has no column: {', '.join(missing)}")
    return index


def _row_is_blank(row: tuple) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _row_values(row: tuple, spec: SheetSpec, index: dict[str, int]) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for column in spec.columns:
        position = index.get(column.field)
        if position is None:
            continue
        raw = row[position] if position < len(row) else None
        value = _convert(column, raw)
        if column.required and not column.allow_blank and not value:
            raise _InvalidRowError(f"blank {column.headers[0]!r}")
        values[column.field] = value
    return values


def _raw_key(row: tuple, spec: SheetSpec, index: dict[str, int]) -> tuple[str, ...] | None:
    """The business key of a row that failed validation, or None when a key
    cell is missing or blank (and may not be)."""
    columns = {column.field: column for column in spec.columns}
    parts: list[str] = []
    for field in spec.key_fields:
        position = index.get(field)
        if position is None:
            return None
        value = _text(row[position]) if position < len(row) else ""
        if not value and not columns[field].allow_blank:
            return None
        parts.append(value)
    return tuple(parts)


def _parse_sheet(workbook: Any, spec: SheetSpec) -> tuple[list[ParsedRow], list[SkippedRow]]:
    if spec.sheet_name not in workbook.sheetnames:
        raise MissingSheetError(
            f"Sheet {spec.sheet_name!r} not found in workbook (available: {workbook.sheetnames!r})"
        )
    rows = workbook[spec.sheet_name].iter_rows(min_row=1, values_only=True)
    header = next(rows, None)
    if header is None:
        return [], []
    index = _header_index(header, spec)

    parsed: list[ParsedRow] = []
    skipped: list[SkippedRow] = []
    seen: set[tuple[str, ...]] = set()
    for row_number, row in enumerate(rows, start=2):
        if _row_is_blank(row):
            continue
        try:
            values = _row_values(row, spec, index)
        except _InvalidRowError as exc:
            skipped.append(
                SkippedRow(spec.sheet_name, row_number, str(exc), _raw_key(row, spec, index))
            )
            continue
        key = tuple(values[field] for field in spec.key_fields)
        if key in seen:
            skipped.append(SkippedRow(spec.sheet_name, row_number, "duplicate key", key))
            continue
        seen.add(key)
        parsed.append(ParsedRow(row_number, key, values))
    return parsed, skipped


def parse_workbook(data: bytes) -> ParsedWorkbook:
    """Parse `data` (an `.xlsx` file's bytes) into the three sheets' rows."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        results = [_parse_sheet(workbook, spec) for spec in SHEETS]
    finally:
        workbook.close()
    (products, skipped_p), (series, skipped_s), (admins, skipped_a) = results
    return ParsedWorkbook(
        products=tuple(products),
        series=tuple(series),
        admins=tuple(admins),
        skipped=tuple(skipped_p + skipped_s + skipped_a),
    )
