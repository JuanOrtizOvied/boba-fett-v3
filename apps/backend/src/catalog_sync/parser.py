"""Excel workbook parser: bytes -> raw per-row field values, keyed by
`FIELD_MAPPING`'s DB field names (`openspec/changes/catalog-sharepoint-sync`
— design.md ADR-2, ADR-5; spec scenarios SYNC-05, SYNC-06, SYNC-14 to
SYNC-18, SYNC-27).

Pure module: reads the workbook bytes it's given, no network, no database.
Values land here as raw text (numeric cells coerced to text — SYNC-18);
composite parsing (`composite.parse_composite`) and official-list/currency
matching (`normalize.py`) run later, once the official lists a full match
needs (including the live `manager`/`administrator` tables) are available —
this module only extracts what the sheet itself contains.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import load_workbook

from catalog_sync.config import (
    CONTROL_FIELD_TYPE,
    DELETE_FLAG_TRUE_VALUES,
    FIELD_MAPPING,
    KEY_FIELD_TYPE,
    SHEET_NAME,
    FieldMapping,
)
from catalog_sync.normalize import normalize_key


class MissingSheetError(ValueError):
    """Raised when the configured sheet is not present in the workbook
    (SYNC-17). The caller must abort the run without writing anything."""


@dataclass(frozen=True)
class ParsedRow:
    """One data row, with `codigo` and the `Eliminar` flag pulled out (both
    drive the diff engine's control flow) and every other mapped column as
    raw text, keyed by DB field name."""

    row_number: int
    codigo: str
    fields: dict[str, str]
    delete_flag: bool


@dataclass(frozen=True)
class ParseReport:
    """Rows the parser skipped, for the sync's run log (never raises — a
    blank or duplicate `codigo` is an expected, recoverable condition)."""

    skipped_blank_codigo: tuple[int, ...] = ()
    skipped_duplicate_codigo: tuple[tuple[int, str], ...] = ()


@dataclass(frozen=True)
class ParseResult:
    rows: tuple[ParsedRow, ...]
    report: ParseReport


def _cell_to_text(value: object) -> str:
    """Coerce a cell value to text (SYNC-18: `Flujos = 0` -> `"0"`). Whole
    floats drop the trailing `.0` since Excel stores plain integers typed
    into a numeric cell as floats."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _parse_delete_flag(value: object) -> bool:
    """`Eliminar` is true for the boolean `True` and for the (case
    -insensitive) strings/number in `DELETE_FLAG_TRUE_VALUES`. Blank,
    `FALSE`, and anything else mean no delete signal (design.md ADR-5)."""
    if isinstance(value, bool):
        return value
    return _cell_to_text(value).upper() in DELETE_FLAG_TRUE_VALUES


def _build_header_index(
    header_row: tuple, field_mapping: dict[str, FieldMapping]
) -> dict[int, str]:
    """Map 1-indexed column position -> DB field, by matching each header
    cell's normalized text against the normalized `FIELD_MAPPING` keys
    (SYNC-14: column order doesn't matter). A header with no match is
    ignored (SYNC-15); a `FIELD_MAPPING` entry with no matching header
    (e.g. a missing `Eliminar` column, SYNC-16) simply has no index."""
    normalized_lookup = {
        normalize_key(header): entry["db_field"] for header, entry in field_mapping.items()
    }

    index: dict[int, str] = {}
    for col, cell in enumerate(header_row, start=1):
        header_text = cell.value
        if not header_text:
            continue
        db_field = normalized_lookup.get(normalize_key(str(header_text)))
        if db_field is not None:
            index[col] = db_field
    return index


def _row_is_blank(row: tuple) -> bool:
    return all(cell.value is None or str(cell.value).strip() == "" for cell in row)


def parse_workbook(
    data: bytes,
    *,
    sheet_name: str = SHEET_NAME,
    field_mapping: dict[str, FieldMapping] = FIELD_MAPPING,
) -> ParseResult:
    """Parse `data` (an `.xlsx` file's bytes) into `ParsedRow`s plus a
    report of skipped rows.

    Raises `MissingSheetError` if `sheet_name` isn't in the workbook
    (SYNC-17) — the caller must abort the whole run without writing
    anything when this happens.
    """
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise MissingSheetError(
                f"Sheet {sheet_name!r} not found in workbook "
                f"(available: {workbook.sheetnames!r})"
            )
        sheet = workbook[sheet_name]

        db_field_by_type = {entry["db_field"]: entry["type"] for entry in field_mapping.values()}

        rows_iter = sheet.iter_rows(min_row=1)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return ParseResult(rows=(), report=ParseReport())

        header_index = _build_header_index(header_row, field_mapping)
        codigo_column = next(
            (col for col, db_field in header_index.items() if db_field == "codigo"), None
        )
        eliminar_column = next(
            (col for col, db_field in header_index.items() if db_field == "is_deleted"), None
        )

        parsed_rows: list[ParsedRow] = []
        skipped_blank: list[int] = []
        skipped_duplicate: list[tuple[int, str]] = []
        seen_codigos: set[str] = set()

        for row_number, row in enumerate(rows_iter, start=2):
            if _row_is_blank(row):
                continue

            raw_codigo = row[codigo_column - 1].value if codigo_column else None
            codigo = _cell_to_text(raw_codigo).strip().upper()

            if not codigo:
                skipped_blank.append(row_number)
                continue
            if codigo in seen_codigos:
                skipped_duplicate.append((row_number, codigo))
                continue
            seen_codigos.add(codigo)

            delete_flag = (
                _parse_delete_flag(row[eliminar_column - 1].value) if eliminar_column else False
            )

            fields: dict[str, str] = {}
            for col, db_field in header_index.items():
                field_type = db_field_by_type[db_field]
                if field_type in (KEY_FIELD_TYPE, CONTROL_FIELD_TYPE):
                    continue
                fields[db_field] = _cell_to_text(row[col - 1].value)

            parsed_rows.append(
                ParsedRow(
                    row_number=row_number,
                    codigo=codigo,
                    fields=fields,
                    delete_flag=delete_flag,
                )
            )

        return ParseResult(
            rows=tuple(parsed_rows),
            report=ParseReport(
                skipped_blank_codigo=tuple(skipped_blank),
                skipped_duplicate_codigo=tuple(skipped_duplicate),
            ),
        )
    finally:
        workbook.close()
