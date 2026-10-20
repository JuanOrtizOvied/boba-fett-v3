"""Tests for the client-template export (`db.portfolio_template_export`)."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from db.models import AssetAllocation, Product
from db.portfolio_template_export import (
    NOTE_SHEET_TITLE,
    build_portfolio_template_workbook,
    template_export_filename,
)

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent
    / "src"
    / "db"
    / "templates"
    / "portafolio_final_template.xlsx"
)


def _alloc(name: str, pct: float = 100) -> list[AssetAllocation]:
    return [AssetAllocation(name=name, percentage=pct)]


def _product(asset_class: str, underlying: str | None, amount: float = 1000, **kw) -> Product:
    defaults = dict(
        user_id="usr_test",
        name="Some Fund",
        amount=amount,
        asset_class=_alloc(asset_class),
        underlying=_alloc(underlying) if underlying else [],
    )
    defaults.update(kw)
    return Product(**defaults)


def _build(products: list[Product]):
    return load_workbook(build_portfolio_template_workbook(products))


def _note_rows(wb) -> dict[str, list]:
    """Data rows of the note sheet keyed by class label (below the header row)."""
    ws = wb[NOTE_SHEET_TITLE]
    header_row = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Clase de activo"
    )
    rows = {}
    for r in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label:
            rows[label] = [ws.cell(row=r, column=c).value for c in range(1, 6)]
    return rows


def test_fidelity_merges_fills_and_formulas_preserved():
    template = load_workbook(TEMPLATE_PATH).worksheets[0]
    wb = _build([_product("mercados_publicos", "Renta Variable US Large Cap")])
    ws = wb.worksheets[0]

    assert len(template.merged_cells.ranges) == 28
    assert {str(r) for r in ws.merged_cells.ranges} == {
        str(r) for r in template.merged_cells.ranges
    }
    for ref in ("C6", "B43", "B87"):
        assert ws[ref].fill.start_color.rgb == template[ref].fill.start_color.rgb
    for ref in ("D20", "D94"):
        assert isinstance(ws[ref].value, str) and ws[ref].value.startswith("=")
        assert ws[ref].value == template[ref].value


def test_clean_underlying_lands_in_leaf_as_fraction():
    wb = _build([_product("mercados_publicos", "Renta Variable US Large Cap", amount=500)])
    assert wb.worksheets[0]["D59"].value == pytest.approx(1.0)


def test_fraction_of_total_across_products():
    wb = _build(
        [
            _product("mercados_publicos", "Renta Variable US Large Cap", amount=250),
            _product("mercados_privados", "Deuda Privada", amount=750),
        ]
    )
    ws = wb.worksheets[0]
    assert ws["D59"].value == pytest.approx(0.25)
    assert ws["D28"].value == pytest.approx(0.75)


def test_alias_rv_prefix_lands_in_leaf():
    wb = _build([_product("mercados_publicos", "RV US Large Cap")])
    assert wb.worksheets[0]["D59"].value == pytest.approx(1.0)


def test_alias_rf_prefix_lands_in_leaf():
    wb = _build([_product("mercados_publicos", "RF US Treasuries")])
    assert wb.worksheets[0]["D65"].value == pytest.approx(1.0)


def test_single_leaf_group_matches_by_group_and_leaf_with_accents():
    wb = _build([_product("mercados_privados", "deuda privada")])
    assert wb.worksheets[0]["D28"].value == pytest.approx(1.0)


def test_composite_group_leaf_match_for_club_deals():
    wb = _build([_product("club_deals", "Real Estate Extranjero")])
    assert wb.worksheets[0]["D46"].value == pytest.approx(1.0)


def test_cash_routes_whole_class_to_d89_even_without_underlying():
    wb = _build([_product("cash_y_equivalentes", None, amount=400)])
    assert wb.worksheets[0]["D89"].value == pytest.approx(1.0)
    _, mapped, unmapped, total, _ = _note_rows(wb)["Cash y equivalentes"]
    assert (mapped, unmapped, total) == (pytest.approx(1), pytest.approx(0), pytest.approx(1))


def test_unmapped_underlying_not_written_and_reported_with_reconciliation():
    products = [
        _product("mercados_publicos", "Renta Variable US Large Cap", amount=600),
        _product("mercados_publicos", "Mystery Fund XYZ", amount=400, name="Mystery Fund XYZ"),
    ]
    wb = _build(products)
    ws = wb.worksheets[0]
    assert ws["D59"].value == pytest.approx(0.6)
    leaf_cells = [ws.cell(row=r, column=4).value for r in range(1, ws.max_row + 1)]
    assert not any(isinstance(v, float) and v == pytest.approx(0.4) for v in leaf_cells)

    rows = _note_rows(wb)
    label, mapped, unmapped, total, detail = rows["Mercados públicos"]
    assert mapped == pytest.approx(0.6)
    assert unmapped == pytest.approx(0.4)
    assert total == pytest.approx(1.0)
    assert mapped + unmapped == pytest.approx(total)
    assert "Mystery Fund XYZ" in detail


def test_empty_underlying_is_fully_unmapped():
    wb = _build([_product("club_deals", None, amount=100)])
    rows = _note_rows(wb)
    _, mapped, unmapped, total, _ = rows["Club deals"]
    assert (mapped, unmapped, total) == (pytest.approx(0), pytest.approx(1), pytest.approx(1))


def test_partial_underlying_remainder_is_unmapped():
    product = Product(
        user_id="usr_test",
        name="Some Fund",
        amount=1000,
        asset_class=_alloc("mercados_privados"),
        underlying=[
            AssetAllocation(name="Private Equity", percentage=70),
            AssetAllocation(name="Something Else", percentage=30),
        ],
    )
    wb = _build([product])
    assert wb.worksheets[0]["D30"].value == pytest.approx(0.7)
    _, mapped, unmapped, total, _ = _note_rows(wb)["Mercados privados"]
    assert mapped == pytest.approx(0.7)
    assert unmapped == pytest.approx(0.3)
    assert total == pytest.approx(1.0)


def test_legacy_asset_class_alias_resolves():
    wb = _build([_product("privados", "Hedge Funds")])
    assert wb.worksheets[0]["D36"].value == pytest.approx(1.0)


def test_empty_portfolio_returns_blank_actual_and_empty_note():
    wb = _build([])
    ws = wb.worksheets[0]
    for ref in ("D17", "D28", "D59", "D77", "D89"):
        assert ws[ref].value is None
    assert NOTE_SHEET_TITLE in wb.sheetnames
    assert _note_rows(wb) == {}


def test_filename_shape():
    assert re.fullmatch(r"portafolio-sabbi-\d{4}-\d{2}-\d{2}\.xlsx", template_export_filename())


def test_on_disk_template_is_not_mutated():
    build_portfolio_template_workbook([_product("inversiones_directas", "Residencial")])
    assert load_workbook(TEMPLATE_PATH).worksheets[0]["D17"].value is None
