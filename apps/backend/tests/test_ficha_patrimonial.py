"""Tests for the ficha patrimonial Excel parser (`db.ficha_patrimonial`).

Covers `normalize_return_rate()`, `TIPO_ACTIVO_MAP` / `map_asset_class()`,
and `parse_ficha_excel()` against fixture workbooks built with openpyxl
(`build_ficha_workbook` in `conftest.py`) — `sdd/admin-ficha-patrimonial/spec`
"Excel Parsing — Sabbi Sheet", "Return Rate Normalization", and "Asset Class
Mapping".
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

# --- normalize_return_rate() -------------------------------------------------


@pytest.mark.parametrize(
    "raw_value, expected",
    [
        ("7.5%", "7.5"),
        (0.08, "8.0"),
        ("8", "8.0"),
        ("", ""),
        (None, ""),
        ("  10%  ", "10.0"),
    ],
)
def test_normalize_return_rate(raw_value, expected):
    from db.ficha_patrimonial import normalize_return_rate

    assert normalize_return_rate(raw_value) == expected


def test_normalize_return_rate_returns_empty_for_non_numeric_garbage():
    from db.ficha_patrimonial import normalize_return_rate

    assert normalize_return_rate("n/a") == ""


# --- TIPO_ACTIVO_MAP / map_asset_class() -------------------------------------


@pytest.mark.parametrize(
    "tipo_activo, expected",
    [
        ("Acciones en bolsa", "mercados_publicos"),
        ("Fondos mutuos", "mercados_publicos"),
        ("Cuenta ahorros", "cash_y_equivalentes"),
        ("Cuenta corriente", "cash_y_equivalentes"),
        ("Deposito plazo fijo", "cash_y_equivalentes"),
        (
            "Inversiones alternativas (private equity/venture capital/etc)",
            "mercados_privados",
        ),
    ],
)
def test_map_asset_class_known_values(tipo_activo, expected):
    from db.ficha_patrimonial import map_asset_class

    assert map_asset_class(tipo_activo) == expected


def test_map_asset_class_is_case_and_accent_insensitive():
    from db.ficha_patrimonial import map_asset_class

    assert map_asset_class("CUENTA AHORROS") == "cash_y_equivalentes"
    assert map_asset_class("Depósito Plazo Fijo") == "cash_y_equivalentes"


@pytest.mark.parametrize("tipo_activo", ["Otro", "Criptomonedas", "", None])
def test_map_asset_class_unknown_value_returns_none(tipo_activo):
    from db.ficha_patrimonial import map_asset_class

    assert map_asset_class(tipo_activo) is None


# --- parse_ficha_excel() -----------------------------------------------------


def test_parse_ficha_excel_extracts_client_info(build_ficha_workbook):
    from db.ficha_patrimonial import parse_ficha_excel

    file_bytes = build_ficha_workbook(
        products=[{"name": "Fondo A", "amount": 1000}],
        client_name="Juan Perez",
        client_email="juan.perez@example.com",
        client_phone="+51 999 999 999",
    )

    parsed = parse_ficha_excel(file_bytes)

    assert parsed.client.name == "Juan Perez"
    assert parsed.client.email == "juan.perez@example.com"
    assert parsed.client.phone == "+51 999 999 999"


def test_parse_ficha_excel_returns_exactly_expected_row_range(build_ficha_workbook):
    from db.ficha_patrimonial import parse_ficha_excel

    products = [
        {
            "name": f"Producto {i}",
            "tipo_activo": "Fondos mutuos",
            "currency": "USD",
            "pertenencia": "Titular",
            "amount": 1000 * i,
            "return_rate": "7.5%",
        }
        for i in range(1, 16)
    ]
    file_bytes = build_ficha_workbook(products=products, include_total_row=True)

    parsed = parse_ficha_excel(file_bytes)

    assert len(parsed.rows) == 15
    assert parsed.rows[0].excel_row == 8
    assert parsed.rows[-1].excel_row == 22
    assert parsed.rows[0].raw_name == "Producto 1"
    assert parsed.rows[0].mapped_asset_class == "mercados_publicos"
    assert parsed.rows[0].normalized_return_rate == "7.5"


def test_parse_ficha_excel_stops_before_blank_row(build_ficha_workbook):
    from db.ficha_patrimonial import parse_ficha_excel

    products = [
        {"name": "Producto 1", "amount": 1000},
        {"name": "Producto 2", "amount": 2000},
    ]
    file_bytes = build_ficha_workbook(products=products, include_total_row=False)

    parsed = parse_ficha_excel(file_bytes)

    assert len(parsed.rows) == 2


def test_parse_ficha_excel_skips_row_with_missing_amount(build_ficha_workbook):
    from db.ficha_patrimonial import parse_ficha_excel

    products = [
        {"name": "Producto 1", "amount": 1000},
        {"name": "Producto sin monto", "amount": 0},
        {"name": "Producto 3", "amount": 3000},
    ]
    file_bytes = build_ficha_workbook(products=products, include_total_row=True)

    parsed = parse_ficha_excel(file_bytes)

    names = [row.raw_name for row in parsed.rows]
    assert "Producto sin monto" not in names
    assert names == ["Producto 1", "Producto 3"]


def test_parse_ficha_excel_excludes_real_estate_rows():
    from db.ficha_patrimonial import parse_ficha_excel

    wb = Workbook()
    ws = wb.active
    ws.title = "Sabbi"
    ws["D5"] = "Juan Perez"
    ws["M5"] = "juan.perez@example.com"
    ws["O5"] = "+51 999 999 999"

    # Product rows 8-10 (no Total row — table runs directly into real estate)
    for offset, name in enumerate(["Producto 1", "Producto 2", "Producto 3"]):
        row = 8 + offset
        ws[f"C{row}"] = name
        ws[f"M{row}"] = 1000

    # Real estate section starting at row 27 (`REAL_ESTATE_START_ROW`)
    ws["C27"] = "Departamento Miraflores"
    ws["M27"] = 250000

    parsed = parse_ficha_excel(wb_to_bytes(wb))

    names = [row.raw_name for row in parsed.rows]
    assert names == ["Producto 1", "Producto 2", "Producto 3"]
    assert "Departamento Miraflores" not in names


def test_parse_ficha_excel_raises_on_missing_sabbi_sheet():
    from db.ficha_patrimonial import FichaParseError, parse_ficha_excel

    wb = Workbook()
    wb.active.title = "OtherSheet"

    with pytest.raises(FichaParseError, match="Sabbi"):
        parse_ficha_excel(wb_to_bytes(wb))


def test_parse_ficha_excel_raises_on_corrupt_file():
    from db.ficha_patrimonial import FichaParseError, parse_ficha_excel

    with pytest.raises(FichaParseError):
        parse_ficha_excel(b"not a real xlsx file")


def wb_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
